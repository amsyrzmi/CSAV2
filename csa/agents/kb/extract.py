"""Upload validation and text extraction for knowledge base files (LG13).

Console uploads are PDF, DOCX or CSV files stored under
``<KB_UPLOAD_ROOT>/<client_id>/``. Every path is resolved (following symlinks)
and must stay inside that tenant directory; size, extension and file signature
are checked before any parser runs.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Sequence
from pathlib import Path
from typing import Literal
from uuid import UUID

from pydantic import BaseModel, ConfigDict, Field

from csa.agents.kb.chunking import DEFAULT_MAX_CHARS, chunk_text

FileKind = Literal["pdf", "docx", "csv"]

_EXTENSIONS: dict[str, FileKind] = {".pdf": "pdf", ".docx": "docx", ".csv": "csv"}
MAX_PDF_PAGES = 500
MAX_CSV_ROWS = 5000
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
_QUESTION_COLUMNS = ("question", "questions", "q", "soalan", "pertanyaan")
_ANSWER_COLUMNS = ("answer", "answers", "a", "jawapan")


class KbFileError(ValueError):
    """A single uploaded file is unusable; the message is safe to show staff."""


class ExtractedChunk(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    content: str = Field(min_length=1)
    heading: str | None = None
    question: str | None = None


class ExtractedDocument(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    kind: Literal["doc", "faq"]
    file_kind: FileKind
    chunks: list[ExtractedChunk] = Field(min_length=1)


def file_kind_for(name: str) -> FileKind:
    suffix = Path(name).suffix.lower()
    if suffix == ".doc":
        raise KbFileError("legacy .doc files are not supported; save the file as .docx")
    kind = _EXTENSIONS.get(suffix)
    if kind is None:
        raise KbFileError(f"unsupported file type '{suffix or name}'; use PDF, DOCX or CSV")
    return kind


def resolve_upload_path(upload_root: Path, client_id: UUID, path: str) -> Path:
    """Resolve ``path`` and require it to be a regular file in the tenant directory."""
    if not path or "\x00" in path:
        raise KbFileError("file path is empty or invalid")
    tenant_dir = (upload_root / str(client_id)).resolve()
    candidate = Path(path)
    if not candidate.is_absolute():
        candidate = tenant_dir / candidate
    try:
        resolved = candidate.resolve(strict=True)
    except (FileNotFoundError, RuntimeError, OSError) as exc:
        raise KbFileError("file not found in the client upload directory") from exc
    if not resolved.is_relative_to(tenant_dir):
        raise KbFileError("file is outside the client upload directory")
    if not resolved.is_file():
        raise KbFileError("path is not a regular file")
    return resolved


def read_upload(path: Path, max_bytes: int) -> bytes:
    with path.open("rb") as fh:
        data = fh.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise KbFileError(f"file exceeds the {max_bytes // (1024 * 1024)} MB limit")
    if not data:
        raise KbFileError("file is empty")
    return data


def extract_document(
    data: bytes,
    file_kind: FileKind,
    *,
    title: str,
    max_chars: int = DEFAULT_MAX_CHARS,
) -> ExtractedDocument:
    if file_kind == "pdf":
        return _extract_pdf(data, title, max_chars)
    if file_kind == "docx":
        return _extract_docx(data, title, max_chars)
    return _extract_csv(data, title, max_chars)


def _sections_to_doc(
    sections: Sequence[tuple[str | None, str]], file_kind: FileKind, max_chars: int
) -> ExtractedDocument:
    chunks = [
        ExtractedChunk(content=piece, heading=heading)
        for heading, text in sections
        for piece in chunk_text(text, max_chars)
    ]
    if not chunks:
        raise KbFileError("no readable text found in file")
    return ExtractedDocument(kind="doc", file_kind=file_kind, chunks=chunks)


def _extract_pdf(data: bytes, title: str, max_chars: int) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    if not data.startswith(b"%PDF-"):
        raise KbFileError("file is not a valid PDF")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise KbFileError("password-protected PDFs are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise KbFileError(f"PDF has more than {MAX_PDF_PAGES} pages")
        sections = [
            (f"{title} (page {number})", text)
            for number, page in enumerate(reader.pages, start=1)
            if (text := (page.extract_text() or "").strip())
        ]
    except KbFileError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise KbFileError(f"PDF could not be read: {type(exc).__name__}") from exc
    if not sections:
        raise KbFileError("PDF has no extractable text (scanned images need OCR first)")
    return _sections_to_doc(sections, "pdf", max_chars)


def _extract_docx(data: bytes, title: str, max_chars: int) -> ExtractedDocument:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    if not data.startswith(b"PK\x03\x04"):
        raise KbFileError("file is not a valid DOCX")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if "word/document.xml" not in zf.namelist():
                raise KbFileError("file is not a valid DOCX")
            if sum(info.file_size for info in zf.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise KbFileError("DOCX expands beyond the allowed size")
        document = Document(io.BytesIO(data))
    except KbFileError:
        raise
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError, ValueError) as exc:
        raise KbFileError(f"DOCX could not be read: {type(exc).__name__}") from exc

    sections: list[tuple[str | None, str]] = []
    heading: str | None = title
    buffer: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            if buffer:
                sections.append((heading, "\n\n".join(buffer)))
                buffer = []
            heading = text[:255]
        else:
            buffer.append(text)
    if buffer:
        sections.append((heading, "\n\n".join(buffer)))
    for table in document.tables:
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            for row in table.rows
        ]
        table_text = "\n\n".join(r for r in rows if r)
        if table_text:
            sections.append((f"{title} (table)", table_text))
    if not sections:
        raise KbFileError("DOCX has no readable text")
    return _sections_to_doc(sections, "docx", max_chars)


def _decode_csv(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise KbFileError("file is not a text CSV")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on Windows commonly exports CSV as cp1252.
        try:
            return data.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise KbFileError("CSV must be UTF-8 or Windows-1252 encoded") from exc


def _find_column(headers: list[str], names: tuple[str, ...]) -> str | None:
    for header in headers:
        if re.sub(r"[^a-z]", "", header.lower()) in names:
            return header
    return None


def _extract_csv(data: bytes, title: str, max_chars: int) -> ExtractedDocument:
    text = _decode_csv(data)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    headers = [h for h in (reader.fieldnames or []) if h and h.strip()]
    if not headers:
        raise KbFileError("CSV has no header row")

    rows: list[dict[str, str]] = []
    try:
        for index, row in enumerate(reader):
            if index >= MAX_CSV_ROWS:
                raise KbFileError(f"CSV has more than {MAX_CSV_ROWS} rows")
            cleaned = {h: (row.get(h) or "").strip() for h in headers}
            if any(cleaned.values()):
                rows.append(cleaned)
    except csv.Error as exc:
        raise KbFileError(f"CSV could not be parsed: {exc}") from exc
    if not rows:
        raise KbFileError("CSV has no data rows")

    q_col = _find_column(headers, _QUESTION_COLUMNS)
    a_col = _find_column(headers, _ANSWER_COLUMNS)
    if q_col and a_col:
        chunks = [
            ExtractedChunk(
                content=f"Q: {row[q_col]}\nA: {row[a_col]}"[: max_chars * 4],
                heading=title,
                question=row[q_col][:1000],
            )
            for row in rows
            if row[q_col] and row[a_col]
        ]
        if not chunks:
            raise KbFileError("CSV FAQ has no rows with both a question and an answer")
        return ExtractedDocument(kind="faq", file_kind="csv", chunks=chunks)

    lines = ["; ".join(f"{h}: {row[h]}" for h in headers if row[h]) for row in rows]
    return _sections_to_doc([(title, "\n\n".join(lines))], "csv", max_chars)
