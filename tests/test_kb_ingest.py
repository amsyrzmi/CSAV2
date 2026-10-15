"""Console upload ingestion for the Knowledge Base agent (LG13).

Covers file validation and extraction (PDF, DOCX, CSV), tenant directory
confinement, strict TEI embedding, and the ``kb_ingest_jobs`` queue processor
against a disposable Postgres with RLS. All data is synthetic.
"""

from __future__ import annotations

import asyncio
import io
import json
import os
import uuid
from collections.abc import AsyncIterator
from pathlib import Path
from typing import Any

import httpx
import psycopg
import pytest
from csa.agents.base import AgentContext
from csa.agents.kb import DefaultKbEmbedder, KbAgent, KbEmbeddingError, KbInput
from csa.agents.kb.extract import (
    KbFileError,
    extract_document,
    file_kind_for,
    read_upload,
    resolve_upload_path,
)
from csa.agents.kb.ingest_jobs import KbIngestConfig, KbIngestProcessor, parse_job_files
from csa.contracts.callbacks import PendingActionStore
from csa.core import db
from docx import Document
from psycopg.rows import dict_row

from tests.conftest import owner_dsn_for, requires_docker

# --- synthetic file builders -------------------------------------------------


def make_pdf(pages: list[str]) -> bytes:
    """Build a minimal valid text PDF with one Helvetica text line per page."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(len(pages)))
    font_id = 3 + len(pages) * 2
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for i, text in enumerate(pages):
        content_id = 4 + i * 2
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
            ).encode()
        )
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx(heading: str, paragraphs: list[str]) -> bytes:
    document = Document()
    document.add_heading(heading, level=1)
    for para in paragraphs:
        document.add_paragraph(para)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


FAQ_CSV = (
    b"Question,Answer\n"
    b"What are your opening hours?,We open 9am to 6pm Monday to Saturday.\n"
    b"Do you deliver to Sabah?,Yes, delivery to Sabah takes 5 to 7 working days.\n"
)


# --- extraction and validation (no database) --------------------------------


def test_pdf_extraction_keeps_page_headings() -> None:
    doc = extract_document(
        make_pdf(["Warranty covers twelve months.", "Returns accepted within 14 days."]),
        "pdf",
        title="Policy",
    )
    assert doc.kind == "doc"
    assert [c.heading for c in doc.chunks] == ["Policy (page 1)", "Policy (page 2)"]
    assert "twelve months" in doc.chunks[0].content


def test_docx_extraction_uses_heading_styles() -> None:
    doc = extract_document(
        make_docx("Shipping", ["Standard delivery takes three days.", "Express is next day."]),
        "docx",
        title="Guide",
    )
    assert doc.chunks[0].heading == "Shipping"
    assert "Express is next day." in doc.chunks[0].content


def test_csv_with_question_answer_columns_becomes_faq() -> None:
    doc = extract_document(FAQ_CSV, "csv", title="FAQ")
    assert doc.kind == "faq"
    assert [c.question for c in doc.chunks] == [
        "What are your opening hours?",
        "Do you deliver to Sabah?",
    ]
    assert doc.chunks[1].content.startswith("Q: Do you deliver to Sabah?\nA: Yes")


def test_csv_without_faq_columns_is_row_text() -> None:
    doc = extract_document(b"sku;name;price\nA1;Blue Mug;12.50\n", "csv", title="Catalogue")
    assert doc.kind == "doc"
    assert doc.chunks[0].content == "sku: A1; name: Blue Mug; price: 12.50"


@pytest.mark.parametrize(
    ("data", "kind", "message"),
    [
        (b"not a pdf", "pdf", "not a valid PDF"),
        (b"plain text", "docx", "not a valid DOCX"),
        (b"Question,Answer\n", "csv", "no data rows"),
        (b"\x00\x01binary", "csv", "not a text CSV"),
        (make_pdf([""]), "pdf", "no extractable text"),
    ],
)
def test_invalid_files_raise_safe_errors(data: bytes, kind: Any, message: str) -> None:
    with pytest.raises(KbFileError, match=message):
        extract_document(data, kind, title="x")


def test_file_type_allowlist() -> None:
    assert file_kind_for("Price List.PDF") == "pdf"
    with pytest.raises(KbFileError, match=r"\.doc files are not supported"):
        file_kind_for("old.doc")
    with pytest.raises(KbFileError, match="unsupported file type"):
        file_kind_for("script.exe")


def test_upload_path_is_confined_to_the_tenant_directory(tmp_path: Path) -> None:
    client_a, client_b = uuid.uuid4(), uuid.uuid4()
    (tmp_path / str(client_a)).mkdir()
    (tmp_path / str(client_b)).mkdir()
    own = tmp_path / str(client_a) / "faq.csv"
    own.write_bytes(FAQ_CSV)
    other = tmp_path / str(client_b) / "secret.csv"
    other.write_bytes(FAQ_CSV)
    link = tmp_path / str(client_a) / "link.csv"
    link.symlink_to(other)

    assert resolve_upload_path(tmp_path, client_a, "faq.csv") == own.resolve()
    assert resolve_upload_path(tmp_path, client_a, str(own)) == own.resolve()
    for bad in (f"../{client_b}/secret.csv", str(other), "link.csv"):
        with pytest.raises(KbFileError, match="outside the client upload directory"):
            resolve_upload_path(tmp_path, client_a, bad)
    with pytest.raises(KbFileError, match="not found"):
        resolve_upload_path(tmp_path, client_a, "missing.pdf")
    with pytest.raises(KbFileError, match="not a regular file"):
        resolve_upload_path(tmp_path, client_a, ".")


def test_read_upload_enforces_size_and_emptiness(tmp_path: Path) -> None:
    big = tmp_path / "big.csv"
    big.write_bytes(b"x" * 2048)
    with pytest.raises(KbFileError, match="exceeds"):
        read_upload(big, 1024)
    empty = tmp_path / "empty.csv"
    empty.write_bytes(b"")
    with pytest.raises(KbFileError, match="empty"):
        read_upload(empty, 1024)


def test_parse_job_files_accepts_paths_and_objects() -> None:
    files = parse_job_files(["a.pdf", {"path": "b.docx", "name": "Guide.docx", "x": 1}])
    assert [(f.path, f.name) for f in files] == [("a.pdf", None), ("b.docx", "Guide.docx")]
    with pytest.raises(KbFileError):
        parse_job_files({"path": "a.pdf"})


# --- strict embedder ----------------------------------------------------------


async def test_tei_failures_raise_instead_of_falling_back() -> None:
    def down(request: httpx.Request) -> httpx.Response:
        return httpx.Response(503)

    def wrong_dims(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json=[[0.1, 0.2]])

    for handler, message in ((down, "TEI request failed"), (wrong_dims, "384-dimensional")):
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            embedder = DefaultKbEmbedder("http://tei.test", http_client=client)
            assert embedder.model_name == "tei-384"
            with pytest.raises(KbEmbeddingError, match=message):
                await embedder.embed_batch(["hello"])


async def test_tei_batches_large_inputs() -> None:
    sizes: list[int] = []

    def ok(request: httpx.Request) -> httpx.Response:
        inputs = json.loads(request.content)["inputs"]
        sizes.append(len(inputs))
        return httpx.Response(200, json=[[0.0] * 384 for _ in inputs])

    async with httpx.AsyncClient(transport=httpx.MockTransport(ok)) as client:
        vectors = await DefaultKbEmbedder("http://tei.test", http_client=client).embed_batch(
            ["t"] * 70
        )
    assert len(vectors) == 70 and sizes == [32, 32, 6]


async def test_tei_applies_query_and_passage_prefixes() -> None:
    seen: list[list[str]] = []

    def ok(request: httpx.Request) -> httpx.Response:
        inputs = json.loads(request.content)["inputs"]
        seen.append(inputs)
        return httpx.Response(200, json=[[0.0] * 384 for _ in inputs])

    async with httpx.AsyncClient(transport=httpx.MockTransport(ok)) as client:
        embedder = DefaultKbEmbedder(
            "http://tei.test",
            query_prefix="query: ",
            passage_prefix="passage: ",
            http_client=client,
        )
        await embedder.embed("waktu operasi")
        await embedder.embed_batch(["Kedai buka 9 pagi"])
    assert seen == [["query: waktu operasi"], ["passage: Kedai buka 9 pagi"]]


def test_prod_requires_tei_and_ingest_requires_upload_root(monkeypatch: Any) -> None:
    from csa.core.settings import Settings

    base = {
        "DATABASE_URL": "postgresql://u:p@localhost/db",
        "LG_DATABASE_URL": "postgresql://u:p@localhost/db",
        "GROQ_API_KEY": "synthetic",
        "PUBLIC_BASE_URL": "https://example.test",
        "CREDENTIALS_KEY": "base64:" + "A" * 43 + "=",
    }
    for key in list(os.environ):
        if key.startswith(("KB_", "TEI_")):
            monkeypatch.delenv(key)
    prod = Settings(ENV="prod", **base)  # type: ignore[arg-type]
    with pytest.raises(KbEmbeddingError, match="TEI_ENDPOINT is required"):
        DefaultKbEmbedder.from_settings(prod)
    with pytest.raises(ValueError, match="requires KB_UPLOAD_ROOT"):
        Settings(ENV="dev", KB_INGEST_ENABLED=True, **base)  # type: ignore[arg-type]
    with pytest.raises(ValueError, match="requires TEI_ENDPOINT"):
        Settings(ENV="prod", KB_INGEST_ENABLED=True, KB_UPLOAD_ROOT="/tmp", **base)  # type: ignore[arg-type]


# --- queue processor against Postgres + RLS -----------------------------------


class FlakyEmbedder(DefaultKbEmbedder):
    def __init__(self) -> None:
        super().__init__()
        self.fail = False

    async def embed_batch(self, texts: list[str]) -> list[list[float]]:
        if self.fail:
            raise KbEmbeddingError("TEI request failed: ConnectError")
        return await super().embed_batch(texts)


class Tenant:
    def __init__(self, owner_dsn: str, upload_root: Path) -> None:
        self.owner_dsn = owner_dsn
        self.id = uuid.uuid4()
        self.dir = upload_root / str(self.id)
        self.dir.mkdir(parents=True)
        with psycopg.connect(owner_dsn, autocommit=True) as conn:
            conn.execute(
                "INSERT INTO clients (id, slug, name, platform_version, status) "
                "VALUES (%s,%s,'Synthetic KB','lg13','active')",
                (self.id, f"kb-{self.id}"),
            )

    def put(self, name: str, data: bytes) -> str:
        (self.dir / name).write_bytes(data)
        return name

    def enqueue(self, files: list[Any], meta: dict[str, Any] | None = None) -> uuid.UUID:
        with psycopg.connect(self.owner_dsn, autocommit=True) as conn:
            row = conn.execute(
                "INSERT INTO kb_ingest_jobs (client_id, source, files, meta, requested_by) "
                "VALUES (%s,'upload',%s,%s,'console-test') RETURNING id",
                (self.id, json.dumps(files), json.dumps(meta or {})),
            ).fetchone()
        assert row is not None
        return row[0]

    def job(self, job_id: uuid.UUID) -> dict[str, Any]:
        with psycopg.connect(self.owner_dsn) as conn:
            row = (
                conn.cursor(row_factory=dict_row)
                .execute(
                    "SELECT status, error, attempts, meta, locked_at > now() AS backing_off "
                    "FROM kb_ingest_jobs WHERE id=%s",
                    (job_id,),
                )
                .fetchone()
            )
        assert row is not None
        return dict(row)

    def docs(self) -> list[dict[str, Any]]:
        with psycopg.connect(self.owner_dsn) as conn:
            return [
                dict(r)
                for r in conn.cursor(row_factory=dict_row)
                .execute(
                    "SELECT d.id, d.title, d.kind, d.version, d.is_active, d.file_name, "
                    "d.embed_model, d.chunk_count, "
                    "(SELECT count(*) FROM kb_chunks c WHERE c.document_id=d.id) AS chunks "
                    "FROM kb_documents d WHERE d.client_id=%s ORDER BY d.created_at",
                    (self.id,),
                )
                .fetchall()
            ]


class KbWorld:
    def __init__(self, owner_dsn: str, owner_pool: Any, app_pool: Any, root: Path) -> None:
        self.owner_dsn = owner_dsn
        self.app_pool = app_pool
        self.root = root
        self.embedder = FlakyEmbedder()
        self.processor = KbIngestProcessor(
            owner_pool,
            app_pool,
            embedder=self.embedder,
            config=KbIngestConfig(upload_root=root, max_attempts=2, retry_backoff_seconds=60),
        )

    def tenant(self) -> Tenant:
        return Tenant(self.owner_dsn, self.root)

    async def run(self, tenant: Tenant) -> str | None:
        job = await self.processor.claim(tenant.id)
        return None if job is None else await self.processor.process(job)

    async def ask(self, tenant: Tenant, question: str) -> Any:
        ctx = AgentContext(
            client_id=tenant.id,
            contact_id=uuid.uuid4(),
            trace_id=uuid.uuid4(),
            conversation_id=None,
            idempotency_key=f"kbq-{uuid.uuid4().hex[:8]}",
            channel="web_console",
            pool=self.app_pool,
            pending_actions=PendingActionStore(self.app_pool),
        )
        return await KbAgent(embedder=self.embedder).run(
            ctx, KbInput(action="query", query=question, threshold=0.3)
        )


@pytest.fixture
async def kb(
    migrated_dsn: str, pg_container: object, app_dsn: str, tmp_path: Path
) -> AsyncIterator[KbWorld]:
    owner = owner_dsn_for(pg_container)
    owner_pool = db.build_pool(owner, max_size=2)
    app_pool = db.build_pool(app_dsn, max_size=4)
    await owner_pool.open(wait=True)
    await app_pool.open(wait=True)
    try:
        yield KbWorld(owner, owner_pool, app_pool, tmp_path)
    finally:
        await app_pool.close()
        await owner_pool.close()


@requires_docker
async def test_three_file_upload_is_ingested_and_searchable(kb: KbWorld) -> None:
    tenant = kb.tenant()
    files = [
        tenant.put("policy.pdf", make_pdf(["Warranty claims require the original receipt."])),
        {
            "path": tenant.put("3f2a.bin.docx", make_docx("Delivery", ["Parcels ship daily."])),
            "name": "Delivery Guide.docx",
        },
        tenant.put("faq.csv", FAQ_CSV),
    ]
    job_id = tenant.enqueue(files, {"category": "support", "language": "ms"})

    assert await kb.run(tenant) == "done"
    job = tenant.job(job_id)
    assert job["status"] == "done" and job["error"] is None and job["attempts"] == 1
    assert [r["status"] for r in job["meta"]["results"]] == ["ingested"] * 3

    docs = tenant.docs()
    assert {(d["title"], d["kind"], d["file_name"]) for d in docs} == {
        ("policy", "doc", "policy.pdf"),
        ("Delivery Guide", "doc", "Delivery Guide.docx"),
        ("faq", "faq", "faq.csv"),
    }
    assert all(d["embed_model"] == "hash-384" and d["chunks"] == d["chunk_count"] for d in docs)

    found = await kb.ask(tenant, "warranty claims original receipt")
    assert found.ok and "original receipt" in found.data.facts[0].text
    assert await kb.run(tenant) is None


@requires_docker
async def test_other_tenants_cannot_see_or_claim_uploads(kb: KbWorld) -> None:
    owner, other = kb.tenant(), kb.tenant()
    owner.enqueue([owner.put("faq.csv", FAQ_CSV)])
    assert await kb.run(other) is None
    assert await kb.run(owner) == "done"

    miss = await kb.ask(other, "What are your opening hours?")
    assert miss.ok and miss.data.facts[0].id == "kb_gap_recorded"
    hit = await kb.ask(owner, "What are your opening hours?")
    assert hit.data.facts[0].id.startswith("kb_chunk_")


@requires_docker
async def test_more_than_three_files_rejects_the_job(kb: KbWorld) -> None:
    tenant = kb.tenant()
    job_id = tenant.enqueue([tenant.put(f"f{i}.csv", FAQ_CSV) for i in range(4)])
    assert await kb.run(tenant) == "failed"
    job = tenant.job(job_id)
    assert "maximum is 3" in job["error"]
    assert tenant.docs() == []


@requires_docker
async def test_bad_files_fail_individually_as_partial(kb: KbWorld) -> None:
    tenant = kb.tenant()
    sibling = kb.tenant()
    sibling.put("private.csv", FAQ_CSV)
    job_id = tenant.enqueue(
        [
            tenant.put("faq.csv", FAQ_CSV),
            f"../{sibling.id}/private.csv",
            tenant.put("fake.pdf", b"MZ not really a pdf"),
        ]
    )
    assert await kb.run(tenant) == "partial"
    results = tenant.job(job_id)["meta"]["results"]
    assert [r["status"] for r in results] == ["ingested", "failed", "failed"]
    assert "outside the client upload directory" in results[1]["error"]
    assert "not a valid PDF" in results[2]["error"]
    assert len(tenant.docs()) == 1


@requires_docker
async def test_reupload_is_idempotent_and_changed_file_creates_new_version(kb: KbWorld) -> None:
    tenant = kb.tenant()
    tenant.enqueue([tenant.put("faq.csv", FAQ_CSV)])
    assert await kb.run(tenant) == "done"
    again = tenant.enqueue(["faq.csv"])
    assert await kb.run(tenant) == "done"
    assert tenant.job(again)["meta"]["results"][0]["status"] == "duplicate"
    assert len(tenant.docs()) == 1

    tenant.put("faq.csv", FAQ_CSV + b"Is parking free?,Yes for two hours.\n")
    tenant.enqueue(["faq.csv"])
    assert await kb.run(tenant) == "done"
    docs = tenant.docs()
    assert [(d["version"], d["is_active"], d["chunks"]) for d in docs] == [
        (1, False, 0),
        (2, True, 3),
    ]


@requires_docker
async def test_embedding_outage_requeues_with_backoff_then_fails(kb: KbWorld) -> None:
    tenant = kb.tenant()
    job_id = tenant.enqueue([tenant.put("faq.csv", FAQ_CSV)])
    kb.embedder.fail = True

    assert await kb.run(tenant) == "queued"
    job = tenant.job(job_id)
    assert job["status"] == "queued" and job["backing_off"] and "TEI" in job["error"]
    assert await kb.run(tenant) is None  # still backing off
    assert tenant.docs() == []

    with psycopg.connect(kb.owner_dsn, autocommit=True) as conn:
        conn.execute("UPDATE kb_ingest_jobs SET locked_at = now() WHERE id=%s", (job_id,))
    assert await kb.run(tenant) == "failed"
    assert tenant.job(job_id)["attempts"] == 2


@requires_docker
async def test_concurrent_claims_take_a_job_once(kb: KbWorld) -> None:
    tenant = kb.tenant()
    tenant.enqueue([tenant.put("faq.csv", FAQ_CSV)])
    claims = await asyncio.gather(*(kb.processor.claim(tenant.id) for _ in range(4)))
    assert sum(c is not None for c in claims) == 1


@requires_docker
async def test_abandoned_processing_job_is_requeued(kb: KbWorld) -> None:
    tenant = kb.tenant()
    job_id = tenant.enqueue([tenant.put("faq.csv", FAQ_CSV)])
    assert await kb.processor.claim(tenant.id) is not None  # simulated crash: never finished
    with psycopg.connect(kb.owner_dsn, autocommit=True) as conn:
        conn.execute(
            "UPDATE kb_ingest_jobs SET locked_at = now() - interval '1 hour' WHERE id=%s",
            (job_id,),
        )
    assert await kb.run(tenant) == "done"
    assert tenant.job(job_id)["attempts"] == 2
